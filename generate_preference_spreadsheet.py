import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def build_preference_spreadsheet(csv_path="teacher_preferences.csv", output_path="Faculty_Teaching_Preferences.xlsx"):
    # 1. Load CSV data
    df = pd.read_csv(csv_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Preferences Overview"
    ws.views.sheetView[0].showGridLines = True

    # 2. Northeastern University Color Palette
    NU_RED = "CC0000"
    HEADER_FILL = PatternFill(start_color=NU_RED, end_color=NU_RED, fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    ZEBRA_FILL = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    TITLE_FONT = Font(name="Segoe UI", size=16, bold=True, color="111111")
    SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="555555")
    DATA_FONT = Font(name="Segoe UI", size=10, color="111111")

    THIN_BORDER = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )

    # 3. Header Block
    ws["A1"] = "📚 Northeastern Economics — Faculty Teaching Preferences"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Generated Spreadsheet Report"
    ws["A2"].font = SUBTITLE_FONT

    start_row = 4

    # 4. Write Data Table Header and Rows
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

        # Format Headers
        if r_idx == start_row:
            for cell in ws[r_idx]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[r_idx].height = 28
        else:
            # Format Data Rows
            fill_color = ZEBRA_FILL if (r_idx % 2 == 0) else WHITE_FILL
            for cell in ws[r_idx]:
                cell.fill = fill_color
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.row_dimensions[r_idx].height = 22

    # 5. Auto-fit Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.row < start_row:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 6. Save Formatted Excel File
    wb.save(output_path)
    print(f"✅ Excel spreadsheet successfully generated: {output_path}")

if __name__ == "__main__":
    build_preference_spreadsheet()